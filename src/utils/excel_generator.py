"""Утилита генерации Excel-файлов (.xlsx) для отчётов и экспорта данных."""

from __future__ import annotations

import io
from collections.abc import Sequence
from datetime import datetime
from typing import Any

from loguru import logger
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from ..data.models.batch import Batch
from ..data.models.product import Product

# --- Стили ---

HEADER_FONT = Font(bold=True, size=11)
HEADER_FILL = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)

TITLE_FONT = Font(bold=True, size=14)
SUBTITLE_FONT = Font(bold=True, size=11)
LABEL_FONT = Font(bold=True)

# --- Колонки экспорта ---

EXPORT_COLUMNS: list[tuple[str, int]] = [
    ("Номер партии", 20),
    ("Дата смены", 15),
    ("Номер смены", 14),
    ("Статус партии", 16),
    ("Код продукции", 22),
    ("Тип продукции", 16),
    ("Статус продукции", 18),
    ("Дата производства", 20),
    ("Доп. данные", 30),
]


def generate_report(report_data: dict[str, Any]) -> bytes:
    """Генерация Excel-отчёта по производству за период.

    Формирует файл с одним листом, содержащим:
    - Заголовок с периодом и рабочим центром
    - Сводку: всего партий, план, факт, % выполнения
    - Таблицу разбивки по статусам

    Args:
        report_data: Словарь из ``AnalyticsService.get_production_report()``.
            Ожидаемые ключи: ``date_from``, ``date_to``, ``work_center_id``,
            ``total_batches``, ``batches_by_status``, ``total_planned``,
            ``total_actual``, ``completion_percent``.

    Returns:
        Содержимое .xlsx файла в байтах.
    """
    wb = Workbook()
    ws: Worksheet = wb.active
    ws.title = "Производственный отчёт"

    current_row = 1

    # --- Заголовок ---
    ws.cell(row=current_row, column=1, value="Производственный отчёт").font = TITLE_FONT
    current_row += 1

    date_from = report_data.get("date_from", "")
    date_to = report_data.get("date_to", "")
    if isinstance(date_from, datetime):
        date_from = date_from.strftime("%d.%m.%Y")
    if isinstance(date_to, datetime):
        date_to = date_to.strftime("%d.%m.%Y")

    ws.cell(row=current_row, column=1, value=f"Период: {date_from} — {date_to}")
    current_row += 1

    wc_id = report_data.get("work_center_id")
    wc_label = f"Рабочий центр: {wc_id}" if wc_id is not None else "Все рабочие центры"
    ws.cell(row=current_row, column=1, value=wc_label)
    current_row += 2

    # --- Сводка ---
    ws.cell(row=current_row, column=1, value="Сводка").font = SUBTITLE_FONT
    current_row += 1

    summary_items: list[tuple[str, object]] = [
        ("Всего партий", report_data.get("total_batches", 0)),
        ("План (кол-во)", report_data.get("total_planned", 0)),
        ("Факт (кол-во)", report_data.get("total_actual", 0)),
        ("Выполнение, %", report_data.get("completion_percent", 0)),
    ]

    for label, value in summary_items:
        ws.cell(row=current_row, column=1, value=label).font = LABEL_FONT
        ws.cell(row=current_row, column=2, value=str(value))
        current_row += 1

    current_row += 1

    # --- Таблица по статусам ---
    ws.cell(row=current_row, column=1, value="Партии по статусам").font = SUBTITLE_FONT
    current_row += 1

    for col_idx, header_text in enumerate(["Статус", "Количество"], start=1):
        cell = ws.cell(row=current_row, column=col_idx, value=header_text)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGNMENT
    current_row += 1

    batches_by_status: dict[str, int] = report_data.get("batches_by_status", {})
    for status, count in batches_by_status.items():
        ws.cell(row=current_row, column=1, value=status)
        ws.cell(row=current_row, column=2, value=count)
        current_row += 1

    # --- Ширина колонок ---
    ws.column_dimensions["A"].width = 25
    ws.column_dimensions["B"].width = 18

    file_bytes: bytes = _workbook_to_bytes(wb)
    logger.info("Excel-отчёт сгенерирован: {} байт", len(file_bytes))
    return file_bytes


def generate_export(rows: Sequence[Any]) -> bytes:
    """Генерация Excel-файла для экспорта данных (продукция + партии).

    Формирует таблицу с заголовками и данными. Каждая строка — одна
    единица продукции с информацией о партии.

    Args:
        rows: Результат SQL-запроса — последовательность кортежей ``(Product, Batch)``.

    Returns:
        Содержимое .xlsx файла в байтах.
    """
    wb = Workbook()
    ws: Worksheet = wb.active
    ws.title = "Экспорт данных"

    # --- Заголовки ---
    for col_idx, (header_text, width) in enumerate(EXPORT_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header_text)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGNMENT
        col_letter: str = cell.column_letter
        ws.column_dimensions[col_letter].width = width

    # --- Данные ---
    for row_idx, row in enumerate(rows, start=2):
        product: Product = row[0]
        batch: Batch = row[1]

        ws.cell(row=row_idx, column=1, value=batch.number)
        ws.cell(
            row=row_idx,
            column=2,
            value=batch.shift_date.isoformat() if batch.shift_date else "",
        )
        ws.cell(row=row_idx, column=3, value=batch.shift_number)
        ws.cell(row=row_idx, column=4, value=batch.status.value)
        ws.cell(row=row_idx, column=5, value=product.code)
        ws.cell(row=row_idx, column=6, value=product.product_type)
        ws.cell(row=row_idx, column=7, value=product.status.value)
        ws.cell(
            row=row_idx,
            column=8,
            value=product.produced_at.isoformat() if product.produced_at else "",
        )
        ws.cell(row=row_idx, column=9, value=product.extra_data or "")

    file_bytes: bytes = _workbook_to_bytes(wb)
    logger.info("Excel-экспорт сгенерирован: {} строк, {} байт", len(rows), len(file_bytes))
    return file_bytes


def _workbook_to_bytes(wb: Workbook) -> bytes:
    """Сериализация Workbook в байты."""
    buffer = io.BytesIO()
    wb.save(buffer)
    wb.close()
    result: bytes = buffer.getvalue()
    return result
