"""Утилита парсинга Excel-файлов (.xlsx) с данными о продукции."""

from __future__ import annotations

import io

from loguru import logger
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

# Колонки, которые обязательно должны присутствовать в заголовке
REQUIRED_COLUMNS: set[str] = {"code", "product_type"}


def parse_products(file_bytes: bytes) -> list[dict[str, str]]:
    """Парсит Excel-файл и возвращает список словарей с данными о продукции.

    Формат возвращаемых данных совместим с выводом ``csv.DictReader``,
    чтобы дальнейшая обработка в ``src/tasks/imports.py`` была единообразной.

    Args:
        file_bytes: Содержимое Excel-файла в байтах (скачано из MinIO).

    Returns:
        Список словарей вида ``{"code": "...", "product_type": "...", "extra_data": "..."}``.
        Пустые строки пропускаются.

    Raises:
        ValueError: Если файл пустой, не содержит листов или отсутствуют
            обязательные колонки.
    """
    wb = load_workbook(filename=io.BytesIO(file_bytes), read_only=True, data_only=True)

    try:
        ws: Worksheet = wb.active
        if ws is None:
            raise ValueError("Excel-файл не содержит активного листа")

        rows = ws.iter_rows(min_row=1)

        # --- Заголовок ---
        header_row = next(rows, None)
        if header_row is None:
            raise ValueError("Excel-файл пуст (нет строки заголовка)")

        headers: list[str] = [
            str(cell.value).strip().lower() if cell.value is not None else "" for cell in header_row
        ]

        missing = REQUIRED_COLUMNS - set(headers)
        if missing:
            raise ValueError(f"Отсутствуют обязательные колонки: {', '.join(sorted(missing))}")

        logger.debug("Excel заголовки: {}", headers)

        # --- Данные ---
        results: list[dict[str, str]] = []

        for row_idx, row in enumerate(rows, start=2):
            values = [cell.value for cell in row]

            # Пропускаем полностью пустые строки
            if all(v is None for v in values):
                continue

            record: dict[str, str] = {}
            for header, value in zip(headers, values, strict=False):
                if not header:
                    continue
                record[header] = str(value).strip() if value is not None else ""

            # Проверяем обязательные поля
            if not record.get("code"):
                logger.warning("Строка {}: пустой code, пропускаем", row_idx)
                continue

            if not record.get("product_type"):
                logger.warning("Строка {}: пустой product_type, пропускаем", row_idx)
                continue

            results.append(record)

        logger.info("Excel: распознано {} записей", len(results))
        return results
    finally:
        wb.close()
